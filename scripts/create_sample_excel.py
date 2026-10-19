"""
Script to create the sample Excel file with exact sample rows specified in the PRD.
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DATA_ROWS = [
    ("CABL00058", 1, "TZKA8027", "2026_2.2.1.", "TZKA8027-1795051262"),
    ("CABL00059", 1, "TZKA8027", "2026_2.2.1.", "TZKA8027-1795051262"),
    ("DBDS00031", 1, "TZKA8027", "2026_2.2.1.", "TZKA8027-1795051262"),
    ("DBDS00030", 2, "TZKA8027", "2026_2.2.1.", "TZKA8027-1795051262"),
    ("DBDS00029", 1, "TZKA8027", "2026_2.2.1.", "TZKA8027-1795051262"),
    ("DBDS00028", 2, "TZKA8027", "2026_2.2.1.", "TZKA8027-1795051262"),
    ("FOKT00004", 1, "TZKA8027", "2026_2.2.1.", "TZKA8027-1795051262"),
]

HEADERS = [
    "Item Code",
    "Quantity Req Per Site",
    "Site ID",
    "Line",
    "WBS",
]

def generate_sample_workbook(target_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Material Lines"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10)
    border_side = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Append headers
    ws.append(HEADERS)
    for col_num in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border

    # Append data rows ensuring explicit data types (Line as string to preserve trailing dot)
    for row_idx, row_values in enumerate(DATA_ROWS, start=2):
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if col_idx == 2:  # Quantity
                cell.value = int(val)
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = str(val)
                cell.number_format = "@"  # Text format
                cell.alignment = Alignment(horizontal="left")
            cell.font = regular_font
            cell.border = cell_border

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target_path)
    return target_path

if __name__ == "__main__":
    out_file = Path("data/sample_material_lines.xlsx")
    generate_sample_workbook(out_file)
    print(f"Sample workbook generated successfully at: {out_file.resolve()}")
